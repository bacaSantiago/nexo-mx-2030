"""Load CNBV March 2026 access points, CE 2024 internet sales, ENOE 2026-Q2 unemployment, MIP 2018 multipliers, and 2025 PROFECO complaint counts.

Writes drizzle/migrations/0007_supply_sources.sql. Does not read DENUE into the database.
"""

from __future__ import annotations

import csv
import hashlib
import io
import math
import sys
import unicodedata
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from scripts.build_gap_load import Bucket, Z90, emit_inserts, flag_of, q, seg_key

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
OUT = ROOT / "drizzle" / "migrations" / "0007_supply_sources.sql"
VERSION = "2026-10-17"

STATE_NAMES = {
    "aguascalientes": "01",
    "baja california": "02",
    "baja california sur": "03",
    "campeche": "04",
    "coahuila": "05",
    "coahuila de zaragoza": "05",
    "colima": "06",
    "chiapas": "07",
    "chihuahua": "08",
    "ciudad de mexico": "09",
    "distrito federal": "09",
    "durango": "10",
    "guanajuato": "11",
    "guerrero": "12",
    "hidalgo": "13",
    "jalisco": "14",
    "mexico": "15",
    "estado de mexico": "15",
    "michoacan": "16",
    "michoacan de ocampo": "16",
    "morelos": "17",
    "nayarit": "18",
    "nuevo leon": "19",
    "oaxaca": "20",
    "puebla": "21",
    "queretaro": "22",
    "quintana roo": "23",
    "san luis potosi": "24",
    "sinaloa": "25",
    "sonora": "26",
    "tabasco": "27",
    "tamaulipas": "28",
    "tlaxcala": "29",
    "veracruz": "30",
    "veracruz de ignacio de la llave": "30",
    "yucatan": "31",
    "zacatecas": "32",
}

CNBV_INDICATORS = (
    ("cnbv_correspondents", "Corresponsales", "corresponsales"),
    ("cnbv_branches", "Sucursales", "sucursales"),
    ("cnbv_atms", "Cajeros automáticos", "cajeros"),
    ("cnbv_pos_terminals", "Terminales punto de venta", "tpv"),
)

CE_CHANNELS = (
    (5, "Páginas web de empresas proveedoras", "proveedoras"),
    (6, "Páginas web de intermediarios", "intermediarios"),
    (7, "Redes sociales", "redes sociales"),
    (8, "Correo electrónico", "correo"),
    (9, "Intercambio electrónico de datos", "intercambio"),
    (10, "Otra modalidad informática", "otra modalidad"),
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def fold(text: str) -> str:
    normalized = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return " ".join(normalized.lower().split())


def num(value) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, str):
        value = value.replace(",", "").strip()
        if not value:
            return None
    number = float(value)
    if not math.isfinite(number):
        return None
    return number


def fact_tuple(metric, year, source, universe, geo_level, geo, segment, estimate, unit, se, lo, hi, cv, precision, method, note, url) -> str:
    def sql_num(value, digits):
        return "null" if value is None else f"{value:.{digits}f}"

    conf = "null" if lo is None else "0.90"
    return (
        f"({q(metric)},{year},(select id from public.source_document where source_code={q(source)}),"
        f"{q(universe)},{q(geo_level)},{q(geo)},{q(__import__('json').dumps(segment, ensure_ascii=False))}::jsonb,{q(seg_key(segment))},"
        f"{estimate},{q(unit)},{sql_num(se, 4)},{sql_num(lo, 4)},{sql_num(hi, 4)},{conf},{sql_num(cv, 4)},{q(precision)},'observed',"
        f"(select id from public.dim_method where method_code={q(method)}),{q(VERSION)},null,{q(url)},{q(note)})"
    )


def published(metric, year, source, universe, geo_level, geo, segment, estimate, unit, note, url) -> str:
    return fact_tuple(metric, year, source, universe, geo_level, geo, segment, estimate, unit, None, None, None, None, "high", "official_published_aggregate", note, url)


def read_cnbv_sheet(ws, needle: str, code_col: int, count_cols: dict[str, int], extra_col: int | None):
    totals = None
    header = None
    parsed = {}
    for index, row in enumerate(ws.iter_rows(values_only=True)):
        if index == 1:
            totals = row
        if header is None:
            if any(isinstance(cell, str) and needle in cell for cell in row):
                header = row
            continue
        if row[code_col] is None:
            break
        code = str(int(row[code_col])).zfill(5 if code_col == 1 else 2)
        counts = {key: num(row[col]) for key, col in count_cols.items()}
        if any(value is not None and value < 0 for value in counts.values()):
            raise SystemExit(f"Negative CNBV count in {code}")
        extra = num(row[extra_col]) if extra_col is not None else " ".join(str(row[5]).split())
        parsed[code] = (extra, counts)
    if header is None or totals is None:
        raise SystemExit(f"No header {needle} in {ws.title}")
    return totals, header, parsed


def cnbv_rows(path: Path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    mun_totals, mun_header, municipalities = read_cnbv_sheet(
        wb["BD Infraestructura Mun"], "Clave Municipio", 1,
        {"corresponsales": 12, "sucursales": 13, "cajeros": 14, "tpv": 15}, None,
    )
    edo_totals, edo_header, state_rows = read_cnbv_sheet(
        wb["BD Infraestructura Edo"], "Número de municipios", 2,
        {"corresponsales": 11, "sucursales": 12, "cajeros": 13, "tpv": 14}, 7,
    )
    wb.close()
    if "Sucursales" not in str(mun_header[13]) or "Sucursales" not in str(edo_header[12]):
        raise SystemExit("CNBV column layout is not the March 2026 infrastructure block")
    if "Sucursales" not in str(mun_header[20]) or "Sucursales" not in str(edo_header[19]):
        raise SystemExit("CNBV sector blocks moved; refusing to guess which sucursales column is the total")

    national = {
        "corresponsales": int(mun_totals[12]),
        "sucursales": int(mun_totals[13]),
        "cajeros": int(mun_totals[14]),
        "tpv": int(mun_totals[15]),
    }
    if int(edo_totals[12]) != national["sucursales"]:
        raise SystemExit("State and municipal sheets publish different national branch totals")
    unidentified = municipalities.pop("99999", None)
    state_rows.pop("99", None)
    municipalities = {
        code: (name, code[:2], counts)
        for code, (name, counts) in municipalities.items()
        if code[:2] in {f"{i:02d}" for i in range(1, 33)}
    }
    states = {code: row for code, row in state_rows.items() if code in {f"{i:02d}" for i in range(1, 33)}}
    if len(municipalities) < 2400 or len(states) != 32:
        raise SystemExit(f"Unexpected CNBV coverage: {len(municipalities)} municipalities, {len(states)} states")
    if unidentified:
        national["no_identificado"] = unidentified[1]
    return national, municipalities, states


def reconciliation(national, municipalities, states) -> str:
    parts = []
    for key, label in (("sucursales", "sucursales"), ("corresponsales", "corresponsales"), ("cajeros", "cajeros"), ("tpv", "TPV")):
        mun_sum = sum(counts[key] or 0 for _name, _parent, counts in municipalities.values())
        state_sum = sum(counts[key] or 0 for _adults, counts in states.values())
        extra = (national.get("no_identificado") or {}).get(key) or 0
        parts.append(
            f"{label}: encabezado {national[key]:,}, entidades {state_sum:,.0f}, municipios {mun_sum:,.0f}, sin clave {extra:,.0f}".replace(",", " ")
        )
    return "Totales del bloque Todos los sectores, cierre a marzo 2026. La fila sin clave municipal no se carga como geografía. " + ". ".join(parts) + "."


def load_cnbv(path: Path):
    national, municipalities, states = cnbv_rows(path)
    note_counts = reconciliation(national, municipalities, states)
    url = "https://www.cnbv.gob.mx/Inclusi%C3%B3n/Paginas/Bases-de-Datos.aspx"
    rate_note = (
        "Cociente de la cifra publicada y la población adulta de la misma fila, por 10 mil adultos, hoja BD Infraestructura Edo. "
        "No es una tasa muestral y no trae intervalo. "
        + note_counts
    )
    facts = []
    supply = []
    adults_national = sum(adults or 0 for adults, _counts in states.values())
    for code, label, key in CNBV_INDICATORS:
        supply.append(("national", "00", code, label, national[key]))
        if adults_national > 0:
            rate = national[key] / adults_national * 10000
            facts.append(published(
                f"{code}_per_10k_adults", 2026, "cnbv_inclusion_databases", "Población adulta",
                "national", "00", {}, f"{rate:.2f}", "por 10 mil adultos", rate_note, url,
            ))
        for geo, (adults, counts) in states.items():
            value = counts[key]
            if value is None:
                continue
            supply.append(("state", geo, code, label, value))
            if adults and adults > 0:
                facts.append(published(
                    f"{code}_per_10k_adults", 2026, "cnbv_inclusion_databases", "Población adulta",
                    "state", geo, {}, f"{value / adults * 10000:.2f}", "por 10 mil adultos", rate_note, url,
                ))
        for geo, (_name, _parent, counts) in municipalities.items():
            value = counts[key]
            if value is None:
                continue
            supply.append(("municipality", geo, code, label, value))
    geos = [("municipality", geo, name, parent) for geo, (name, parent, _counts) in municipalities.items()]
    print(note_counts)
    print("cnbv", len(municipalities), "municipalities", len(facts), "rates", len(supply), "counts")
    return facts, supply, geos, note_counts


def load_ce(path: Path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    header = None
    data_rows = []
    for index, row in enumerate(ws.iter_rows(values_only=True)):
        if header is None:
            if row and row[0] == "Entidad federativa":
                header = row
            continue
        if not row[0] or not str(row[0]).startswith("00"):
            continue
        data_rows.append(row)
    wb.close()
    if header is None or "ventas realizadas por internet" not in str(header[4]):
        raise SystemExit("CE header is not the internet-sales tabulation")
    for index, label, needle in CE_CHANNELS:
        if needle not in fold(str(header[index])):
            raise SystemExit(f"CE column {index} is not {label}")
    total = next(row for row in data_rows if row[1] == "Total" and row[2] == "Total")
    income = num(total[3])
    internet = num(total[4])
    if income is None or internet is None or not (40_000_000 < income < 50_000_000) or not (8_000_000 < internet < 9_000_000):
        raise SystemExit("CE national totals are not the expected 2023 million-peso figures")
    url = "https://www.inegi.org.mx/programas/ce/2024/"
    note = (
        "Censos Económicos 2024, datos 2023, tabulado TEMACE_1. Millones de pesos. "
        "El archivo solo trae la entidad 00 Estados Unidos Mexicanos. "
        "Los estratos y las actividades parciales no se suman al total."
    )
    channel_note = note + " Las seis modalidades del total nacional suman el ingreso por ventas por internet."
    universe = "Unidades económicas"
    facts = [
        published("ce_total_income_mxn_millions", 2023, "censos_economicos_2024", universe, "national", "00", {}, f"{income:.3f}", "millones de pesos", note, url),
        published("ce_internet_sales_mxn_millions", 2023, "censos_economicos_2024", universe, "national", "00", {}, f"{internet:.3f}", "millones de pesos", note, url),
        published("ce_internet_sales_share_pct", 2023, "censos_economicos_2024", universe, "national", "00", {}, f"{internet / income * 100:.1f}", "percent", note + " Es el cociente de las dos cifras publicadas, no un porcentaje muestral.", url),
    ]
    for index, label, _needle in CE_CHANNELS:
        value = num(total[index])
        if value is None:
            raise SystemExit(f"Missing CE channel {label}")
        facts.append(published(
            "ce_internet_sales_mxn_millions", 2023, "censos_economicos_2024", universe, "national", "00",
            {"sales_channel": label}, f"{value:.3f}", "millones de pesos", channel_note, url,
        ))
    activities = 0
    for row in data_rows:
        if row[2] != "Total" or row[1] == "Total":
            continue
        value = num(row[4])
        if value is None:
            continue
        activity = " ".join(str(row[1]).split())
        facts.append(published(
            "ce_internet_sales_mxn_millions", 2023, "censos_economicos_2024", universe, "national", "00",
            {"economic_activity": activity}, f"{value:.3f}", "millones de pesos", note, url,
        ))
        activities += 1
    print("ce activities", activities, "share", round(internet / income * 100, 1))
    return facts


def load_enoe(path: Path):
    buckets: dict[str, Bucket] = defaultdict(Bucket)
    pea = unemployed = skipped = 0
    with zipfile.ZipFile(path) as archive:
        name = next(item for item in archive.namelist() if item.endswith("sdem_enoe_2026_2t.csv"))
        with archive.open(name) as handle:
            reader = csv.DictReader(io.TextIOWrapper(handle, encoding="latin-1"))
            for row in reader:
                if str(row["r_def"]).zfill(2) != "00" or row["c_res"] not in {"1", "3"}:
                    continue
                try:
                    age = int(row["eda"])
                    weight = float(row["fac_tri"])
                except ValueError:
                    continue
                if age < 15 or age > 98 or weight <= 0 or row["clase1"] != "1":
                    continue
                if row["clase2"] not in {"1", "2"}:
                    skipped += 1
                    continue
                if not row["est_d_tri"] or not row["upm"]:
                    skipped += 1
                    continue
                y = 1.0 if row["clase2"] == "2" else 0.0
                ent = str(int(row["cve_ent"])).zfill(2)
                stratum, psu = row["est_d_tri"], row["upm"]
                buckets["00"].add(stratum, psu, weight, y, 1.0)
                buckets[ent].add(stratum, psu, weight, y, 1.0)
                pea += 1
                unemployed += int(y)
    if skipped > pea * 0.01:
        raise SystemExit(f"ENOE dropped {skipped} PEA rows for missing class or PSU")
    national = buckets["00"].finish()
    if national is None:
        raise SystemExit("ENOE national unemployment has no variance")
    rate = national[0] * 100
    print("enoe national unemployment", round(rate, 2), "sample pea", pea)
    if not 1.5 <= rate <= 6.0:
        raise SystemExit(f"ENOE national rate {rate:.2f} is outside 1.5–6.0; not loading")
    url = "https://www.inegi.org.mx/programas/enoe/15ymas/?ps=Microdatos"
    note = (
        "Tasa de desocupación calculada del SDEMT, ENOE 2026 segundo trimestre: desocupados entre PEA. "
        "Entrevista definitiva, residentes habituales o nuevos, 15 años o más. Factor fac_tri. "
        "Error estándar de Taylor con est_d_tri y upm, intervalo al 90 %. "
        "Puede diferir del redondeo de un boletín."
    )
    facts = []
    for geo, bucket in buckets.items():
        stats = bucket.finish()
        if stats is None or stats[2] < 2000:
            continue
        proportion, se, den = stats
        center = proportion * 100
        se_pp = se * 100
        cv = se / proportion * 100 if proportion else None
        if cv is None:
            continue
        precision = flag_of(cv)
        if den < 15000 and precision == "high":
            precision = "low"
        level = "national" if geo == "00" else "state"
        facts.append(fact_tuple(
            "unemployment_rate_pct", 2026, "enoe_2026_2t", "Población económicamente activa de 15 años y más",
            level, geo, {}, f"{center:.2f}", "percent", se_pp, max(0.0, center - Z90 * se_pp), min(100.0, center + Z90 * se_pp),
            cv, precision, "survey_design_estimate", note, url,
        ))
    if len(facts) != 33:
        raise SystemExit(f"Expected national plus 32 ENOE rates, got {len(facts)}")
    return facts


def load_mip(path: Path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = next(row for row in rows if row and row[0] == "Concepto")
    data = [row for row in rows if row and row[0] and str(row[0]).strip()[:3].isdigit()]
    labels = [str(cell).strip() for cell in header[1:] if cell]
    if len(data) != 78 or len(labels) != 78 or [str(row[0]).strip() for row in data] != labels:
        raise SystemExit(f"MIP_60 is not the 78-sector Leontief inverse ({len(data)} rows, {len(labels)} columns)")
    url = "https://www.inegi.org.mx/app/simuladormip18/default.html"
    note = (
        "Multiplicador de producción: suma de la columna de la matriz inversa de Leontief, "
        "matriz simétrica insumo-producto 2018, SCIAN, millones de pesos de base. "
        "No se usó para recalcular scenario_result."
    )
    facts = []
    for index, label in enumerate(labels):
        total = sum(num(row[index + 1]) or 0 for row in data)
        if total <= 0:
            raise SystemExit(f"MIP multiplier for {label} is not positive")
        facts.append(published(
            "mip_output_multiplier", 2018, "mip_2018_simulator", "Economía nacional",
            "national", "00", {"sector": label}, f"{total:.4f}", "multiplicador", note, url,
        ))
    print("mip multipliers", len(facts), "agriculture", facts[0][:80])
    return facts


def load_profeco(path: Path):
    by_state: Counter[str] = Counter()
    by_reason: Counter[str] = Counter()
    national = partial = blank = 0
    latest_2026 = ""
    with path.open(encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        if "motivo_reclamacion" not in (reader.fieldnames or []):
            raise SystemExit(f"{path.name} is not the PROFECO complaint register")
        for row in reader:
            year = str(row.get("anio_creacion") or "").strip()
            if year == "2026":
                partial += 1
                fecha = str(row.get("fecha_ingreso") or "")
                if fecha.startswith("2026"):
                    latest_2026 = max(latest_2026, fecha[:10])
            if year != "2025":
                continue
            national += 1
            state = fold(row.get("estado") or "")
            if not state:
                blank += 1
                continue
            code = STATE_NAMES.get(state)
            if not code:
                raise SystemExit(f"Unknown PROFECO state: {row.get('estado')}")
            by_state[code] += 1
            reason = " ".join(str(row.get("motivo_reclamacion") or "").split())
            if reason:
                by_reason[reason[:160]] += 1
    for code in (f"{i:02d}" for i in range(1, 33)):
        by_state.setdefault(code, 0)
    if blank > national * 0.01 or sum(by_state.values()) != national:
        raise SystemExit(f"PROFECO 2025 coverage failed: {national} complaints, {sum(by_state.values())} assigned, {blank} blank")
    missing = [code for code, count in by_state.items() if count == 0]
    url = "https://www.gob.mx/profeco"
    note = "Conteo de expedientes del padrón abierto de quejas con año de creación 2025. Es un conteo administrativo, no una tasa ni una estimación muestral."
    if missing:
        note += " Sin expedientes de 2025 en este archivo: " + ", ".join(missing) + "."
    partial_note = f"Conteo de expedientes con año de creación 2026 en el mismo archivo. El corte llega a {latest_2026 or 'una fecha no indicada'} y no es el cierre del año."
    universe = "Expedientes de queja"
    facts = [published("profeco_complaints", 2025, "profeco_optional", universe, "national", "00", {}, str(national), "expedientes", note, url)]
    for code, count in sorted(by_state.items()):
        facts.append(published("profeco_complaints", 2025, "profeco_optional", universe, "state", code, {}, str(count), "expedientes", note, url))
    for reason, count in by_reason.most_common(8):
        facts.append(published(
            "profeco_complaints", 2025, "profeco_optional", universe, "national", "00",
            {"claim_reason": reason}, str(count), "expedientes", note + " Un expediente puede no estar en estos ocho motivos.", url,
        ))
    facts.append(published("profeco_complaints", 2026, "profeco_optional", universe, "national", "00", {}, str(partial), "expedientes", partial_note, url))
    print("profeco 2025", national, "2026 partial", partial, "through", latest_2026)
    return facts


def source_update(code: str, **fields) -> str:
    sets = ["updated_at = now()"]
    for key, value in fields.items():
        if key == "modules_used":
            inner = ",".join(q(item) for item in value)
            sets.append(f"modules_used = array[{inner}]::text[]")
        elif key == "retrieved_at":
            sets.append("retrieved_at = now()")
        else:
            sets.append(f"{key} = {q(value)}")
    return f"update public.source_document set {', '.join(sets)} where source_code = {q(code)};"


def main() -> None:
    cnbv_path = next(RAW.glob("Base_de_Datos_de_Inclusion_Financiera_202603.xlsx"))
    ce_path = RAW / "CE_TEMACE_1.xlsx"
    enoe_path = next(RAW.glob("*enoe_2026_2t_csv.zip"))
    mip_path = RAW / "MIP_60.xlsx"
    profeco_path = next(path for path in RAW.glob("*.csv") if path.name != "ACCESO DEM NAL.csv" and path.stat().st_size > 1_000_000)
    facts, supply, geos, cnbv_note = load_cnbv(cnbv_path)
    facts += load_ce(ce_path)
    facts += load_enoe(enoe_path)
    facts += load_mip(mip_path)
    facts += load_profeco(profeco_path)

    sql = ["-- Generated by scripts/build_supply_load.py. Do not edit by hand.", "begin;"]
    sql.append(
        "insert into public.source_document(source_code, source_name, title, year, tier, resource_type, official_url, purpose, bundle_status, status) values "
        "('enoe_2026_2t','ENOE 2026 segundo trimestre','Microdatos SDEMT del segundo trimestre de 2026',2026,'core','CSV',"
        "'https://www.inegi.org.mx/programas/enoe/15ymas/?ps=Microdatos','Ocupación y empleo','microdata loaded','curated_included') "
        "on conflict (source_code) do nothing;"
    )
    sql.append(
        "insert into public.dim_metric(metric_code, metric_label, description, default_unit, universe, universe_type, higher_is_better, status, source_id) "
        "select v.c, v.l, v.d, v.u, v.uni, v.t, v.h, 'published_aggregate', (select id from public.source_document where source_code = v.sc) from (values "
        "('cnbv_correspondents_per_10k_adults','Corresponsales por 10 mil adultos','Cociente de corresponsales y población adulta publicados por la CNBV.','por 10 mil adultos','Población adulta','person',true,'cnbv_inclusion_databases'),"
        "('cnbv_branches_per_10k_adults','Sucursales por 10 mil adultos','Cociente de sucursales y población adulta publicados por la CNBV.','por 10 mil adultos','Población adulta','person',true,'cnbv_inclusion_databases'),"
        "('cnbv_atms_per_10k_adults','Cajeros por 10 mil adultos','Cociente de cajeros automáticos y población adulta publicados por la CNBV.','por 10 mil adultos','Población adulta','person',true,'cnbv_inclusion_databases'),"
        "('cnbv_pos_terminals_per_10k_adults','TPV por 10 mil adultos','Cociente de terminales punto de venta y población adulta publicados por la CNBV.','por 10 mil adultos','Población adulta','person',true,'cnbv_inclusion_databases'),"
        "('ce_total_income_mxn_millions','Ingresos por suministro de bienes y servicios','Ingreso publicado por los Censos Económicos 2024, datos 2023.','millones de pesos','Unidades económicas','economy',null::boolean,'censos_economicos_2024'),"
        "('ce_internet_sales_mxn_millions','Ingresos por ventas por internet','Ventas por internet publicadas por los Censos Económicos 2024, datos 2023.','millones de pesos','Unidades económicas','economy',null::boolean,'censos_economicos_2024'),"
        "('ce_internet_sales_share_pct','Ventas por internet sobre ingresos','Cociente de las dos cifras nacionales publicadas.','percent','Unidades económicas','economy',null::boolean,'censos_economicos_2024'),"
        "('unemployment_rate_pct','Tasa de desocupación','Desocupados entre la PEA, ENOE 2026 segundo trimestre.','percent','Población económicamente activa de 15 años y más','person',false,'enoe_2026_2t'),"
        "('mip_output_multiplier','Multiplicador de producción MIP 2018','Suma de columna de la inversa de Leontief. No alimenta los escenarios.','multiplicador','Economía nacional','economy',null::boolean,'mip_2018_simulator'),"
        "('profeco_complaints','Quejas PROFECO','Conteo de expedientes del padrón abierto.','expedientes','Expedientes de queja','economy',null::boolean,'profeco_optional')"
        ") v(c,l,d,u,uni,t,h,sc) on conflict (metric_code) do nothing;"
    )
    emit_inserts(
        sql,
        "insert into public.fact_estimate (metric_code, reference_year, source_id, universe, geo_level, geo_code, segment_json, segment_key, estimate, unit, standard_error, ci_low, ci_high, confidence_level, cv, precision_flag, observed_or_modeled, method_id, version, source_page, source_url, note) values",
        facts,
        "on conflict (metric_code, reference_year, geo_level, geo_code, segment_key, version) do nothing;",
    )
    emit_inserts(
        sql,
        "insert into public.dim_geography(geo_level, geo_code, geo_name, parent_geo_code) values",
        [f"({q(level)},{q(geo)},{q(name)},{q(parent)})" for level, geo, name, parent in geos],
        "on conflict (geo_level, geo_code) do nothing;",
    )
    sql.append("create temporary table _cnbv (geo_level text, geo_code text, indicator_code text, indicator_label text, estimate numeric) on commit drop;")
    emit_inserts(
        sql,
        "insert into _cnbv (geo_level, geo_code, indicator_code, indicator_label, estimate) values",
        [f"({q(level)},{q(geo)},{q(code)},{q(label)},{int(round(value))})" for level, geo, code, label, value in supply],
        ";",
    )
    sql.append(
        "insert into public.fact_supply (indicator_code, indicator_label, reference_year, source_id, universe, geo_level, geo_code, estimate, unit, precision_flag, observed_or_modeled, method_id, version) "
        "select t.indicator_code, t.indicator_label, 2026, s.id, 'Puntos de acceso del sistema financiero', t.geo_level, t.geo_code, t.estimate, 'puntos', 'high', 'observed', m.id, "
        f"{q(VERSION)} from _cnbv t "
        "join public.source_document s on s.source_code = 'cnbv_inclusion_databases' "
        "join public.dim_method m on m.method_code = 'official_published_aggregate' "
        "where not exists (select 1 from public.fact_supply f where f.indicator_code = t.indicator_code and f.reference_year = 2026 and f.geo_level = t.geo_level and f.geo_code = t.geo_code and f.version = "
        f"{q(VERSION)});"
    )
    sql.append(
        "insert into public.dim_segment(segment_code, dimension, value, label) "
        "select distinct lower(regexp_replace(k||'_'||left(v, 80),'[^a-zA-Z0-9]+','_','g')), k, v, left(v, 180) "
        "from public.fact_estimate, jsonb_each_text(segment_json) e(k,v) "
        "on conflict (segment_code) do nothing;"
    )
    sql.append(source_update(
        "cnbv_inclusion_databases", status="curated_included", institution="CNBV", edition="marzo 2026",
        universe_description="Puntos de acceso: corresponsales, sucursales, cajeros y TPV",
        geographic_coverage="Nacional, 32 entidades y municipios de la base a marzo 2026",
        modules_used=["XLSX", "BD Infraestructura Mun", "BD Infraestructura Edo"],
        sha256=sha256(cnbv_path), retrieved_at=True,
        notes=cnbv_note + " No se cargaron captación ni crédito. ACCESO DEM NAL.csv trae otra población adulta y no se usó.",
    ))
    sql.append(source_update(
        "censos_economicos_2024", status="curated_included", institution="INEGI", edition="2024",
        universe_description="Unidades económicas, ingresos y ventas por internet",
        geographic_coverage="Nacional",
        modules_used=["XLSX", "TEMACE_1"], sha256=sha256(ce_path), retrieved_at=True,
        notes="Datos 2023. El tabulado solo incluye 00 Estados Unidos Mexicanos, así que no hay cifra estatal.",
    ))
    sql.append(source_update(
        "enoe_2026_2t", status="curated_included", institution="INEGI", edition="2026 2T",
        universe_description="Población económicamente activa de 15 años y más",
        geographic_coverage="Nacional y 32 entidades",
        modules_used=["CSV", "SDEMT"], sha256=sha256(enoe_path), retrieved_at=True,
        notes="Tasa de desocupación con fac_tri, est_d_tri y upm. ENOE 2025 sigue sin cargarse.",
    ))
    sql.append(source_update(
        "enoe_2025", notes="Sigue pendiente un trimestre de 2025. El segundo trimestre de 2026 está en la fuente enoe_2026_2t.",
    ))
    sql.append(source_update(
        "mip_2018_simulator", status="curated_included", institution="INEGI", edition="2018",
        universe_description="78 sectores SCIAN de la matriz simétrica",
        geographic_coverage="Nacional",
        modules_used=["XLSX", "MIP_60"], sha256=sha256(mip_path), retrieved_at=True,
        notes="Solo se cargaron los multiplicadores de la inversa de Leontief. MIP_42, MIP_44 y MIP_46 no se volcaron. scenario_result no cambió.",
    ))
    sql.append(source_update(
        "profeco_optional", status="curated_included", institution="PROFECO", edition="2025",
        universe_description="Expedientes de queja",
        geographic_coverage="Nacional y 32 entidades",
        modules_used=["CSV", "expedientes"], sha256=sha256(profeco_path), retrieved_at=True,
        notes="Conteo por entidad del año de creación 2025 y total nacional parcial de 2026. No es una tasa de queja.",
    ))
    sql.append("commit;")
    OUT.write_text("\n".join(sql) + "\n")
    print("wrote", OUT, "facts", len(facts), "supply", len(supply), "bytes", OUT.stat().st_size)


if __name__ == "__main__":
    main()
