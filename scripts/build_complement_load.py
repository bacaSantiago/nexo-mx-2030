"""Build migration 0006.

Loads the EIC 2025 municipal internet tabulados, the EIC population
totals, the ENDUTIH streaming precision table, and the ENDUTIH crosses
that the filters still cannot see: state by schooling, and income by
sex and age. Published points already in the database are left as they are.
"""

from __future__ import annotations

import hashlib
import importlib.util
import re
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("gap", ROOT / "scripts/build_gap_load.py")
gap = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(gap)

SRC = Path("/tmp/nexo-src")
OUT = ROOT / "drizzle/migrations/0006_eic_crosses_sources.sql"
VERSION = "2026-10-16"
Z90 = gap.Z90


def round1(value: float) -> float:
    return float(Decimal(str(value)).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP))


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def flag_of(cv: float) -> str:
    if cv < 15:
        return "high"
    if cv < 30:
        return "moderate"
    return "low"


def fact_tuple(metric, year, source, universe, geo_level, geo, segment, estimate, unit, se, lo, hi, cv, precision, observed, method, note, url) -> str:
    se_sql = "null" if se is None else f"{se:.6f}"
    lo_sql = "null" if lo is None else f"{lo:.6f}"
    hi_sql = "null" if hi is None else f"{hi:.6f}"
    cv_sql = "null" if cv is None else f"{cv:.6f}"
    conf = "null" if lo is None else "0.90"
    return (
        f"({gap.q(metric)},{year},(select id from public.source_document where source_code={gap.q(source)}),"
        f"{gap.q(universe)},{gap.q(geo_level)},{gap.q(geo)},{gap.q(__import__('json').dumps(segment, ensure_ascii=False))}::jsonb,{gap.q(gap.seg_key(segment))},"
        f"{estimate},'{unit}',{se_sql},{lo_sql},{hi_sql},{conf},{cv_sql},{gap.q(precision)},{gap.q(observed)},"
        f"(select id from public.dim_method where method_code={gap.q(method)}),{gap.q(VERSION)},null,{gap.q(url)},{gap.q(note)})"
    )


def direct_crosses(people) -> list[str]:
    buckets: dict[str, gap.Bucket] = defaultdict(gap.Bucket)
    for row in people:
        try:
            age = int(row["EDAD"])
            weight = float(row["FAC_PER"])
        except ValueError:
            continue
        if age < 6 or age > 97 or weight <= 0:
            continue
        ent = row["CVE_ENT"].zfill(2)
        school = gap.reviewed.SCHOOL_LABEL.get(row["NIVEL"].zfill(2))
        if not school:
            continue
        user = row["P7_1"] == "1"
        buyer = 1.0 if user and row["P7_21"] == "1" else 0.0
        sex = gap.SEX.get(row["SEXO"])
        buckets[f"sch|{ent}|{school}"].add(row["EST_DIS"], row["UPM_DIS"], weight, 1.0 if user else 0.0, 1.0)
        if user:
            buckets[f"schbuy|{ent}|{school}"].add(row["EST_DIS"], row["UPM_DIS"], weight, buyer, 1.0)
        if sex:
            buckets[f"schsex|{ent}|{school}|{sex}"].add(row["EST_DIS"], row["UPM_DIS"], weight, 1.0 if user else 0.0, 1.0)
    note = (
        "Estimación directa ENDUTIH 2025 (persona seleccionada, FAC_PER). Este cruce estatal por escolaridad no está en la presentación. "
        "Error estándar por linealización de Taylor con EST_DIS y UPM_DIS. Intervalo al 90 %. "
        "Se omite la celda si el denominador expandido es menor a 2 000 personas."
    )
    url = "https://www.inegi.org.mx/contenidos/programas/endutih/2025/datosabiertos/conjunto_de_datos_endutih_2025_csv.zip"
    rows = []

    def push(name: str, geo: str, segment: dict[str, str], metric: str, universe: str) -> None:
        stats = buckets[name].finish()
        if not stats or stats[2] < gap.MIN_DEN:
            return
        p, se, _den = stats
        center = round1(p * 100)
        se_pp = se * 100
        if center <= 0:
            return
        cv = se_pp / center * 100
        precision = flag_of(cv)
        if _den < 15_000 and precision == "high":
            precision = "low"
        rows.append(fact_tuple(
            metric, 2025, "endutih_2025", universe, "state", geo, segment, f"{center:.1f}", "percent",
            se_pp, max(0.0, center - Z90 * se_pp), min(100.0, center + Z90 * se_pp), cv, precision,
            "observed", "survey_design_estimate", note, url,
        ))

    for ent in [f"{i:02d}" for i in range(1, 33)]:
        for school in gap.reviewed.SCHOOL_LABEL.values():
            push(f"sch|{ent}|{school}", ent, {"schooling": school}, "internet_users_pct", "Población de 6 años y más de la entidad")
            push(f"schbuy|{ent}|{school}", ent, {"schooling": school, "use_type": gap.BUY}, "internet_use_type_pct", "Personas usuarias de internet de la entidad")
            for sex in gap.SEX.values():
                push(f"schsex|{ent}|{school}|{sex}", ent, {"schooling": school, "sex": sex}, "internet_users_pct", "Población de 6 años y más de la entidad")
    print(f"state schooling rows {len(rows)}")
    return rows


def income_crosses(people) -> list[str]:
    assigned, meta = gap.load_enigh_deciles()
    cells = gap.load_enigh_people(assigned, meta)
    buckets: dict[str, gap.Bucket] = defaultdict(gap.Bucket)
    matched = skipped = 0
    series = []
    for row in people:
        try:
            age = int(row["EDAD"])
            weight = float(row["FAC_PER"])
        except ValueError:
            continue
        if age < 6 or age > 97 or weight <= 0 or row["SEXO"] not in gap.SEX or row["TLOC"] not in ("1", "2", "3", "4"):
            continue
        school = gap.ENDUTIH_SCHOOL.get(row["NIVEL"].zfill(2), "ne")
        area = "R" if row["TLOC"] == "4" else "U"
        shares = gap.pi_lookup(cells, row["CVE_ENT"].zfill(2), area, row["SEXO"], gap.age_coarse(age), school)
        if not shares:
            skipped += 1
            continue
        matched += 1
        ent = row["CVE_ENT"].zfill(2)
        user = row["P7_1"] == "1"
        buyer = 1.0 if user and row["P7_21"] == "1" else 0.0
        sex = gap.SEX[row["SEXO"]]
        label = gap.age_label(age)
        for index, share in enumerate(shares):
            if share <= 0:
                continue
            weighted = weight * share
            decile = f"Decil {index + 1}"
            buckets[f"check|00|{decile}"].add(row["EST_DIS"], row["UPM_DIS"], weighted, 1.0 if user else 0.0, 1.0)
            if sex and label:
                buckets[f"sxage|00|{decile}|{sex}|{label}"].add(row["EST_DIS"], row["UPM_DIS"], weighted, 1.0 if user else 0.0, 1.0)
                if user:
                    buckets[f"sxagebuy|00|{decile}|{sex}|{label}"].add(row["EST_DIS"], row["UPM_DIS"], weighted, buyer, 1.0)
            if sex:
                buckets[f"sx|{ent}|{decile}|{sex}"].add(row["EST_DIS"], row["UPM_DIS"], weighted, 1.0 if user else 0.0, 1.0)
                if user:
                    buckets[f"sxbuy|{ent}|{decile}|{sex}"].add(row["EST_DIS"], row["UPM_DIS"], weighted, buyer, 1.0)
    print(f"income match persons {matched} skipped {skipped}")
    for decile in range(1, 11):
        stats = buckets[f"check|00|Decil {decile}"].finish()
        series.append(None if not stats else round1(stats[0] * 100))
    print("income check by decile", series)
    if any(v is None for v in series) or series[-1] < series[0] + 5:
        raise SystemExit(f"Income match did not reproduce the expected gradient: {series}")
    note = (
        "Decil de hogar según el ingreso corriente total de ENIGH 2024 (ing_cor). "
        "Cada persona de ENDUTIH recibe la distribución de deciles de su celda ENIGH. "
        "El error estándar es el de muestreo de ENDUTIH condicionado a esa distribución y no incluye la varianza del emparejamiento. "
        "Por eso la precisión no se marca alta. Cruce de tres filtros calculado con el mismo emparejamiento ya cargado."
    )
    url = "https://www.inegi.org.mx/contenidos/programas/enigh/nc/2024/microdatos/enigh2024_ns_concentradohogar_csv.zip"
    rows = []

    def push(name: str, geo: str, segment: dict[str, str], metric: str, universe: str) -> None:
        stats = buckets[name].finish()
        if not stats or stats[2] < gap.MIN_DEN:
            return
        p, se, _den = stats
        center = round1(p * 100)
        se_pp = se * 100
        if center <= 0:
            return
        cv = se_pp / center * 100
        precision = "moderate" if flag_of(cv) == "high" else flag_of(cv)
        rows.append(fact_tuple(
            metric, 2025, "enigh_2024", universe, "state" if geo != "00" else "national", geo, segment,
            f"{center:.1f}", "percent", se_pp, max(0.0, center - Z90 * se_pp), min(100.0, center + Z90 * se_pp),
            cv, precision, "imputed", "multiple_imputation_income", note, url,
        ))

    for decile in range(1, 11):
        label = f"Decil {decile}"
        for sex in gap.SEX.values():
            for _a, _b, age_name in gap.AGE_LABELS:
                push(f"sxage|00|{label}|{sex}|{age_name}", "00", {"income_decile": label, "sex": sex, "age_group": age_name}, "internet_users_pct", "Población de 6 años y más, decil de ingreso imputado")
                push(f"sxagebuy|00|{label}|{sex}|{age_name}", "00", {"income_decile": label, "sex": sex, "age_group": age_name, "use_type": gap.BUY}, "internet_use_type_pct", "Personas usuarias de internet, decil de ingreso imputado")
            for ent in [f"{i:02d}" for i in range(1, 33)]:
                push(f"sx|{ent}|{label}|{sex}", ent, {"income_decile": label, "sex": sex}, "internet_users_pct", "Población de 6 años y más de la entidad, decil de ingreso imputado")
                push(f"sxbuy|{ent}|{label}|{sex}", ent, {"income_decile": label, "sex": sex, "use_type": gap.BUY}, "internet_use_type_pct", "Personas usuarias de internet de la entidad, decil de ingreso imputado")
    print(f"income three-way rows {len(rows)}")
    return rows


def read_internet() -> tuple[list[tuple], list[tuple]]:
    """Municipal and state internet-in-dwellings rows from the 32 EIC workbooks."""
    municipal = []
    states = []
    z_samples = []
    for code in range(1, 33):
        path = SRC / f"eic2025_vivienda_{code:02d}.xlsx"
        if not path.exists():
            raise SystemExit(f"Missing {path.name}. Download it from https://www.inegi.org.mx/contenidos/programas/eic/2025/tabulados/{path.name}")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb["32"]
        block: dict[tuple[str, str], dict[str, float]] = defaultdict(dict)
        names: dict[tuple[str, str], tuple[str, str]] = {}
        for row in ws.iter_rows(min_row=10, values_only=True):
            if row[2] != "Internet" or row[0] is None or row[1] is None or row[3] is None:
                continue
            key = (str(row[0]), str(row[1]))
            names[key] = (str(row[0]), str(row[1]))
            try:
                block[key][str(row[3])] = float(row[5])
            except (TypeError, ValueError):
                continue
        wb.close()
        for key, stats in block.items():
            if "Valor" not in stats:
                continue
            ent_label, mun_label = names[key]
            ent_match = re.match(r"^(\d{2})\b", ent_label.strip())
            if not ent_match:
                continue
            ent = ent_match.group(1)
            value = stats["Valor"]
            se = stats.get("Error estándar")
            lo = stats.get("Límite inferior de confianza")
            hi = stats.get("Límite superior de confianza")
            cv = stats.get("Coeficiente de variación")
            if se and lo is not None and hi is not None and se > 0:
                z_samples.append(((hi - lo) / 2) / se)
            precision = flag_of(cv) if cv is not None else "pending_microdata_recalculation"
            point = round1(value)
            packed = (point, lo, hi, cv, precision)
            if mun_label.strip() == "Total":
                states.append((ent, *packed))
                continue
            mun_match = re.match(r"^(\d{3})\b", mun_label.strip())
            if not mun_match:
                continue
            geo = ent + mun_match.group(1)
            name = mun_label.strip()[3:].strip(" -")
            municipal.append((geo, name, ent, *packed))
    if len(municipal) < 2000:
        raise SystemExit(f"EIC municipal internet rows look incomplete: {len(municipal)}")
    if z_samples:
        mid = sorted(z_samples)[len(z_samples) // 2]
        print(f"EIC interval z median {mid:.3f} (expected about 1.645)")
        if abs(mid - Z90) > 0.05:
            raise SystemExit("EIC confidence interval is not the 90 percent interval this load records")
    print(f"EIC internet municipalities {len(municipal)} states {len(states)}")
    return municipal, states


def read_population() -> list[tuple]:
    path = SRC / "eic2025_poblacion.xlsx"
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["04"]
    block: dict[str, dict[str, float]] = defaultdict(dict)
    for row in ws.iter_rows(min_row=9, values_only=True):
        if row[0] is None or row[1] is None or row[2] is None:
            continue
        try:
            block[str(row[0])][str(row[1])] = float(row[2])
        except (TypeError, ValueError):
            continue
    wb.close()
    out = []
    for label, stats in block.items():
        if "Valor" not in stats:
            continue
        if label.startswith("Estados Unidos"):
            geo, level = "00", "national"
        else:
            match = re.match(r"^(\d{2})\b", label.strip())
            if not match:
                print("skip population label", label)
                continue
            geo, level = match.group(1), "state"
        value = stats["Valor"]
        se = stats.get("Error estándar")
        lo = stats.get("Límite inferior de confianza")
        hi = stats.get("Límite superior de confianza")
        cv = stats.get("Coeficiente de variación")
        out.append((geo, level, round(value), se, lo, hi, cv, flag_of(cv) if cv is not None else "pending_microdata_recalculation"))
    if not any(row[0] == "00" for row in out) or sum(1 for row in out if row[1] == "state") < 32:
        raise SystemExit(f"EIC population table did not yield the national total and 32 states: {len(out)}")
    print(f"EIC population rows {len(out)} national {next(r[2] for r in out if r[0]=='00')}")
    return out


def read_streaming() -> list[tuple]:
    wb = openpyxl.load_workbook(SRC / "2025_h2ur400.xlsx", data_only=True)
    ws = wb["Indicadores"]
    rows = []
    for row in ws.iter_rows(min_row=14, max_row=16, values_only=True):
        area = str(row[0])
        segment = {} if area == "Nacional" else {"area": area}
        if area not in ("Nacional", "Urbano", "Rural"):
            raise SystemExit(f"Unexpected streaming area {area}")
        estimate, se, cv, lo, hi = (float(row[i]) for i in (7, 8, 9, 10, 11))
        rows.append((segment, round1(estimate), se, lo, hi, cv, flag_of(cv)))
    wb.close()
    print("streaming", [(r[0] or "Nacional", r[1]) for r in rows])
    return rows


def source_update(code: str, **fields) -> str:
    sets = ["updated_at = now()"]
    for key, value in fields.items():
        if key == "modules_used":
            inner = ",".join(gap.q(item) for item in value)
            sets.append(f"modules_used = array[{inner}]::text[]")
        elif key == "retrieved_at":
            sets.append("retrieved_at = now()")
        else:
            sets.append(f"{key} = {gap.q(value)}")
    return f"update public.source_document set {', '.join(sets)} where source_code = {gap.q(code)};"


def main() -> None:
    people = gap.reviewed.load_person_rows()
    inserts = direct_crosses(people) + income_crosses(people)
    municipal, state_internet = read_internet()
    population = read_population()
    streaming = read_streaming()

    pop_note = "Población en viviendas particulares habitadas, Encuesta Intercensal 2025, tabulado por entidad. Cifra, error estándar, intervalo al 90 % y CV publicados por INEGI."
    pop_url = "https://www.inegi.org.mx/contenidos/programas/eic/2025/tabulados/eic2025_poblacion.xlsx"
    for geo, level, value, se, lo, hi, cv, precision in population:
        inserts.append(fact_tuple(
            "eic_population_persons", 2025, "eic_2025_population_tabulation",
            "Población en viviendas particulares habitadas", level, geo, {},
            str(int(value)), "personas", se, lo, hi, cv, precision,
            "observed", "official_published_aggregate", pop_note, pop_url,
        ))
    stream_note = "Hogares que disponen de servicio de streaming, ENDUTIH 2025, tabulado h2ur400. Porcentaje, error estándar, intervalo al 90 % y CV publicados. No es el porcentaje de hogares con internet."
    stream_url = "https://www.inegi.org.mx/contenidos/programas/endutih/2025/tabulados/2025_h2ur400.xlsx"
    for segment, estimate, se, lo, hi, cv, precision in streaming:
        inserts.append(fact_tuple(
            "households_streaming_pct", 2025, "endutih_2025_tabulation_h2ur400",
            "Hogares", "national", "00", segment, f"{estimate:.1f}", "percent",
            se, lo, hi, cv, precision, "observed", "official_published_aggregate", stream_note, stream_url,
        ))

    sql = ["-- Generated by scripts/build_complement_load.py. Do not edit by hand.", "begin;"]
    sql.append(
        "insert into public.dim_metric(metric_code, metric_label, description, default_unit, universe, universe_type, higher_is_better, status, source_id) "
        "select v.c, v.l, v.d, v.u, v.uni, v.t, v.h, 'published_aggregate', (select id from public.source_document where source_code = v.sc) from (values "
        "('eic_population_persons','Población en viviendas particulares habitadas','Total publicado por la Encuesta Intercensal 2025. No es la población de 6 años y más de ENDUTIH.','personas','Población en viviendas particulares habitadas','person',null::boolean,'eic_2025_population_tabulation'),"
        "('households_streaming_pct','Hogares con servicio de streaming','Porcentaje de hogares que disponen de servicio de streaming, tabulado h2ur400. No es conexión a internet.','percent','Hogares','household',null::boolean,'endutih_2025_tabulation_h2ur400')"
        ") v(c,l,d,u,uni,t,h,sc) on conflict (metric_code) do nothing;"
    )
    gap.emit_inserts(
        sql,
        "insert into public.fact_estimate (metric_code, reference_year, source_id, universe, geo_level, geo_code, segment_json, segment_key, estimate, unit, standard_error, ci_low, ci_high, confidence_level, cv, precision_flag, observed_or_modeled, method_id, version, source_page, source_url, note) values",
        inserts,
        "on conflict (metric_code, reference_year, geo_level, geo_code, segment_key, version) do nothing;",
    )
    sql.append(
        "insert into public.dim_geography(geo_level, geo_code, geo_name, parent_geo_code) values\n"
        + ",\n".join(f"('municipality',{gap.q(geo)},{gap.q(name)},{gap.q(parent)})" for geo, name, parent, *_rest in municipal)
        + "\non conflict (geo_level, geo_code) do nothing;"
    )
    sql.append(
        "create temporary table _eic (geo_level text, geo_code text, estimate numeric, ci_low numeric, ci_high numeric, cv numeric, precision_flag text) on commit drop;"
    )
    supply_values = []
    for geo, _name, _parent, estimate, lo, hi, cv, precision in municipal:
        supply_values.append(f"('municipality',{gap.q(geo)},{estimate:.1f},{null_num(lo)},{null_num(hi)},{null_num(cv)},{gap.q(precision)})")
    for geo, estimate, lo, hi, cv, precision in state_internet:
        supply_values.append(f"('state',{gap.q(geo)},{estimate:.1f},{null_num(lo)},{null_num(hi)},{null_num(cv)},{gap.q(precision)})")
    gap.emit_inserts(sql, "insert into _eic (geo_level, geo_code, estimate, ci_low, ci_high, cv, precision_flag) values", supply_values, ";")
    sql.append(
        "insert into public.fact_supply (indicator_code, indicator_label, reference_year, source_id, universe, geo_level, geo_code, estimate, unit, ci_low, ci_high, cv, precision_flag, observed_or_modeled, method_id, version) "
        "select 'eic_dwellings_with_internet_pct', 'Viviendas particulares habitadas con internet', 2025, s.id, "
        "'Viviendas particulares habitadas', t.geo_level, t.geo_code, t.estimate, 'percent', t.ci_low, t.ci_high, t.cv, t.precision_flag, 'observed', m.id, "
        f"{gap.q(VERSION)} from _eic t "
        "join public.source_document s on s.source_code = 'eic_2025_housing_tabulation' "
        "join public.dim_method m on m.method_code = 'official_published_aggregate' "
        "where not exists (select 1 from public.fact_supply f where f.indicator_code = 'eic_dwellings_with_internet_pct' and f.reference_year = 2025 and f.geo_level = t.geo_level and f.geo_code = t.geo_code and f.version = "
        f"{gap.q(VERSION)});"
    )
    sql.append(
        "insert into public.dim_segment(segment_code, dimension, value, label) "
        "select distinct lower(regexp_replace(k||'_'||v,'[^a-zA-Z0-9]+','_','g')), k, v, v "
        "from public.fact_estimate, jsonb_each_text(segment_json) e(k,v) "
        "on conflict (segment_code) do nothing;"
    )

    housing_files = [SRC / f"eic2025_vivienda_{i:02d}.xlsx" for i in range(1, 33)]
    housing_hash = hashlib.sha256()
    for path in housing_files:
        housing_hash.update(bytes.fromhex(sha256(path)))
    sql.append(source_update(
        "eic_2025_housing_tabulation", status="curated_included", institution="INEGI", edition="2025",
        universe_description="Viviendas particulares habitadas", geographic_coverage="Municipio y entidad, 32 entidades",
        modules_used=["XLSX", "Vivienda 32 Bienes y TIC"], sha256=housing_hash.hexdigest(), retrieved_at=True,
        notes="Cargados eic2025_vivienda_01.xlsx a eic2025_vivienda_32.xlsx, hoja 32, bien Internet. El enlace de catálogo apunta al archivo de Nayarit; los 32 siguen el mismo directorio.",
    ))
    sql.append(source_update(
        "eic_2025_population_tabulation", status="curated_included", institution="INEGI", edition="2025",
        universe_description="Población en viviendas particulares habitadas", geographic_coverage="Nacional y 32 entidades",
        modules_used=["XLSX", "Población 04"], sha256=sha256(SRC / "eic2025_poblacion.xlsx"), retrieved_at=True,
        notes="Total de población por entidad. Este archivo no trae el desglose municipal.",
    ))
    sql.append(source_update(
        "eic_2025", status="metadata_included", institution="INEGI", edition="2025",
        universe_description="Viviendas particulares habitadas y su población", geographic_coverage="Nacional, entidad y municipio",
        modules_used=["Program page"], retrieved_at=True,
        notes="Los tabulados usados están en las fichas de vivienda y de población.",
    ))
    if (SRC / "iter_00_cpv2020_csv.zip").exists():
        sql.append(source_update(
            "cpv_2020_iter", status="curated_included", institution="INEGI", edition="2020",
            universe_description="Viviendas particulares habitadas", geographic_coverage="Municipios del país",
            modules_used=["CSV", "ITER"], sha256=sha256(SRC / "iter_00_cpv2020_csv.zip"), retrieved_at=True,
            notes="Fila municipal Total del Municipio. VPH_INTER / TVIVPARHAB.",
        ))
    if (SRC / "enigh2024_ns_concentradohogar_csv.zip").exists():
        sql.append(source_update(
            "enigh_2024", status="curated_included", institution="INEGI", edition="2024",
            universe_description="Hogares de ENIGH 2024 nueva serie, personas de 6 años y más",
            geographic_coverage="Nacional, con celdas por entidad",
            modules_used=["CSV", "concentradohogar", "poblacion"],
            sha256=sha256(SRC / "enigh2024_ns_concentradohogar_csv.zip"), retrieved_at=True,
            notes="El decil usa ing_cor del concentrado de hogar y nivelaprob de la tabla de población. El hash es el del concentrado.",
        ))
    if (SRC / "endutih2025.zip").exists():
        sql.append(source_update(
            "endutih_2025", institution="INEGI", edition="2025",
            geographic_coverage="Nacional y 32 entidades",
            sha256=sha256(SRC / "endutih2025.zip"), retrieved_at=True,
        ))
        sql.append(source_update(
            "endutih_2025_presentation", institution="INEGI", edition="2025",
            geographic_coverage="Nacional y 32 entidades", retrieved_at=True,
        ))
    sql.append(source_update(
        "endutih_2025_tabulation_h2ur400", status="curated_included", institution="INEGI", edition="2025",
        universe_description="Hogares", geographic_coverage="Nacional, urbano y rural",
        modules_used=["XLSX", "Indicadores"], sha256=sha256(SRC / "2025_h2ur400.xlsx"), retrieved_at=True,
        notes="Hogares con servicio de streaming. No sustituye el porcentaje de hogares con internet.",
    ))
    sql.append(
        "insert into public.methodology_note(note_code, title, body, topic, source_id) values "
        "('eic_municipal_v1','Internet en viviendas, EIC 2025',"
        "'El porcentaje municipal es el publicado en la hoja 32 de eic2025_vivienda_01 a _32, bien Internet, columna Dispone. "
        "El intervalo y el CV son los del tabulado, al 90 %. No se modela y no se mezcla con el censo 2020.',"
        "'geography', (select id from public.source_document where source_code='eic_2025_housing_tabulation')) "
        "on conflict (note_code) do nothing;"
    )
    sql.append("commit;")
    OUT.write_text("\n".join(sql) + "\n", encoding="utf-8")
    print(f"wrote {OUT} bytes {OUT.stat().st_size} facts {len(inserts)} supply {len(supply_values)}")


def null_num(value) -> str:
    if value is None:
        return "null"
    return f"{float(value):.6f}"


if __name__ == "__main__":
    main()
